"""Turn HTML, PDF, Word and plain-text announcements into line-oriented text.

Table cells are joined with " | " so label/value rows ("Expiration Date | May 8, 2028")
and due-date tables keep their structure for the rule-based extractor.
"""

import html
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser

BLOCK_TAGS = {
    "address", "article", "aside", "blockquote", "br", "dd", "div", "dl", "dt", "fieldset", "figcaption",
    "figure", "footer", "form", "h1", "h2", "h3", "h4", "h5", "h6", "header", "hr", "li", "main", "ol", "p",
    "pre", "section", "table", "tbody", "thead", "tfoot", "tr", "ul", "caption",
}
SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "nav", "iframe", "button", "select"}
CELL_TAGS = {"td", "th"}
HEADING_TAGS = {"h1", "h2", "h3", "h4"}


@dataclass
class Document:
    text: str
    title: str = ""
    description: str = ""
    headings: list = field(default_factory=list)
    links: list = field(default_factory=list)
    site_name: str = ""


class _Parser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skip = 0
        self.in_title = False
        self.title = ""
        self.meta = {}
        self.headings = []
        self._heading = None
        self.links = []
        self._href = None
        self._cell_count = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in SKIP_TAGS:
            self.skip += 1
            return
        if tag == "title":
            self.in_title = True
        elif tag == "meta":
            key = (attrs.get("name") or attrs.get("property") or "").lower()
            if key in ("description", "og:description", "og:title", "og:site_name", "citation_title", "dc.title"):
                self.meta[key] = attrs.get("content") or ""
        elif tag in HEADING_TAGS:
            self._heading = []
        elif tag == "a":
            self._href = attrs.get("href")
        if tag in CELL_TAGS:
            if self._cell_count:
                self.parts.append(" | ")
            self._cell_count += 1
        elif tag == "tr":
            self._cell_count = 0
            self.parts.append("\n")
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in SKIP_TAGS:
            self.skip = max(0, self.skip - 1)
            return
        if tag == "title":
            self.in_title = False
        elif tag in HEADING_TAGS and self._heading is not None:
            text = " ".join("".join(self._heading).split())
            if text:
                self.headings.append(text)
            self._heading = None
        elif tag == "a":
            self._href = None
        if tag == "tr":
            self._cell_count = 0
        if tag in BLOCK_TAGS or tag == "tr":
            self.parts.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title += data
            return
        if self.skip:
            return
        self.parts.append(data)
        if self._heading is not None:
            self._heading.append(data)
        if self._href and data.strip():
            self.links.append((data.strip(), self._href))


def _tidy(text):
    lines = []
    for raw in text.replace("\r", "\n").split("\n"):
        line = re.sub(r"[ \t   ]+", " ", raw).strip(" |")
        line = re.sub(r"(\s*\|\s*)+", " | ", line).strip()
        if line:
            lines.append(line)
        elif lines and lines[-1] != "":
            lines.append("")
    return "\n".join(lines).strip()


def from_html(raw: str) -> Document:
    parser = _Parser()
    try:
        parser.feed(raw)
        parser.close()
    except Exception:  # noqa: BLE001 - malformed markup still yields partial text
        pass
    title = " ".join(html.unescape(parser.title).split())
    meta_title = parser.meta.get("og:title") or parser.meta.get("citation_title") or parser.meta.get("dc.title") or ""
    return Document(
        text=_tidy("".join(parser.parts)),
        title=meta_title.strip() or title,
        description=(parser.meta.get("description") or parser.meta.get("og:description") or "").strip(),
        headings=parser.headings[:40],
        links=parser.links[:400],
        site_name=(parser.meta.get("og:site_name") or "").strip(),
    )


def from_pdf(data: bytes) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages[:400]:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            continue
    title = ""
    try:
        title = (reader.metadata.title or "") if reader.metadata else ""
    except Exception:  # noqa: BLE001
        pass
    return Document(text=_tidy("\n".join(pages)), title=title.strip())


def from_docx(data: bytes) -> Document:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    title = (d.core_properties.title or "").strip()
    return Document(text=_tidy("\n".join(parts)), title=title)


def from_bytes(data: bytes, content_type: str = "", filename: str = "") -> Document:
    name = filename.lower()
    if content_type == "application/pdf" or name.endswith(".pdf") or data[:5] == b"%PDF-":
        return from_pdf(data)
    if name.endswith(".docx") or "officedocument.wordprocessingml" in content_type:
        return from_docx(data)
    text = data.decode("utf-8", errors="replace")
    if "html" in content_type or name.endswith((".htm", ".html")) or re.search(r"<\s*(html|body|div|p|table)\b", text[:5000], re.I):
        return from_html(text)
    return Document(text=_tidy(text))
