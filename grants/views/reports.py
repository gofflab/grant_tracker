import csv
import io
from decimal import Decimal

from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from ..effort import effort_by_year, effort_summary, project_person_months
from ..models import Application, Person
from .common import json_for_chart


def _default_person(request):
    person = getattr(request.user, "person", None)
    if person:
        return person
    return Person.objects.filter(assignments__role__in=["pi", "mpi"]).distinct().first() or Person.objects.filter(
        assignments__isnull=False
    ).distinct().first()


def effort(request):
    rows = effort_summary()
    people = Person.objects.filter(assignments__isnull=False).distinct()
    selected = None
    if request.GET.get("person"):
        selected = people.filter(pk=request.GET["person"]).first()
    selected = selected or _default_person(request)
    chart = None
    if selected:
        this_year = timezone.localdate().year
        years = list(range(this_year - 1, this_year + 5))
        series = effort_by_year(selected, years)
        active = [s for s in series if not s[1]]
        pending = [s for s in series if s[1]]
        spec_series = []
        slot = 1
        for app, is_pending, values in active[:6]:
            spec_series.append({"name": app.display_title, "data": [float(v) for v in values], "slot": slot})
            slot += 1
        if len(active) > 6:
            other = [sum(float(s[2][i]) for s in active[6:]) for i in range(len(years))]
            spec_series.append({"name": "Other active", "data": other, "slot": "muted"})
        if pending:
            spec_series.append({
                "name": "Pending (if all funded)",
                "data": [sum(float(s[2][i]) for s in pending) for i in range(len(years))],
                "slot": 8 if len(spec_series) < 7 else "muted",
            })
        chart = {
            "spec": json_for_chart({"type": "bar", "stacked": True, "labels": [str(y) for y in years],
                                    "series": spec_series, "format": "pm"}),
            "years": years,
            "rows": series,
        }
    return render(request, "grants/effort.html", {
        "rows": rows, "people": people, "selected": selected, "chart": chart,
    })


# ---------------------------------------------------------------------------
# Current & Pending / Other Support
# ---------------------------------------------------------------------------


def _support_items(person, include_preparing=False):
    statuses = [Application.Status.AWARDED, *Application.PENDING]
    if include_preparing:
        statuses += [Application.Status.DRAFTING, Application.Status.ROUTING]
    apps = (
        Application.objects.filter(personnel__person=person, status__in=statuses)
        .select_related("funder", "award")
        .prefetch_related("personnel")
        .distinct()
    )
    active, pending = [], []
    for app in apps:
        award = getattr(app, "award", None)
        if app.status == Application.Status.AWARDED:
            if award and not award.is_active:
                continue
            bucket = active
        else:
            bucket = pending
        lines = [l for l in app.personnel.all() if l.person_id == person.pk]
        role = lines[0].get_role_display() if lines else ""
        start = (award.start_date if award else None) or app.proposed_start
        end = (award.effective_end if award else None) or app.proposed_end
        pi = app.contact_pi or (person.full_name if app.role in ("pi", "mpi") else "")
        bucket.append({
            "app": app,
            "title": app.title,
            "number": (award.award_number if award else "") or app.sponsor_id or app.internal_id,
            "pi": pi,
            "role": role,
            "source": " / ".join(x for x in [app.funder.name if app.funder else "", app.sponsor_unit] if x),
            "place": app.lead_institution or settings.INSTITUTION_NAME,
            "start": start,
            "end": end,
            "amount": (award.awarded_total if award else None) or app.requested_total,
            "goals": app.major_goals or app.abstract,
            "status_label": "Active" if bucket is active else app.get_status_display(),
            "pm_years": project_person_months(app, person),
        })
    key = lambda i: (i["start"] or timezone.localdate())
    return sorted(active, key=key), sorted(pending, key=key)


def _fmt_money(v):
    return f"${Decimal(v):,.0f}" if v is not None else ""


def _fmt_date(d):
    return d.strftime("%m/%Y") if d else ""


def _pm_text(pm_years):
    return "; ".join(f"Year {p['year']} ({p['start'].year}): {p['pm']:.2f} CM".replace(".00 CM", " CM") for p in pm_years if p["pm"]) or "—"


def other_support(request):
    people = Person.objects.filter(assignments__isnull=False).distinct()
    person = people.filter(pk=request.GET.get("person")).first() if request.GET.get("person") else _default_person(request)
    include_preparing = request.GET.get("preparing") == "1"
    active, pending = _support_items(person, include_preparing) if person else ([], [])
    fmt = request.GET.get("format")
    if person and fmt == "csv":
        return _csv(person, active, pending)
    if person and fmt == "docx":
        return _docx(person, active, pending)
    return render(request, "grants/other_support.html", {
        "people": people, "person": person, "active": active, "pending": pending,
        "include_preparing": include_preparing, "institution": settings.INSTITUTION_NAME,
    })


def _csv(person, active, pending):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Status", "Title", "Project number", "PD/PI", "Role", "Source", "Place of performance",
                "Start", "End", "Total award", "Person months", "Major goals"])
    for item in active + pending:
        w.writerow([item["status_label"], item["title"], item["number"], item["pi"], item["role"], item["source"],
                    item["place"], _fmt_date(item["start"]), _fmt_date(item["end"]), item["amount"] or "",
                    _pm_text(item["pm_years"]), item["goals"]])
    resp = HttpResponse(buf.getvalue(), content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="current-pending-{person.last_name.lower()}-{timezone.localdate():%Y%m%d}.csv"'
    return resp


def _docx(person, active, pending):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("Current and Pending (Other) Support")
    r.bold = True
    r.font.size = Pt(13)
    p = d.add_paragraph()
    p.add_run("Name: ").bold = True
    p.add_run(person.full_name)
    if person.orcid:
        p.add_run("    ORCID: ").bold = True
        p.add_run(person.orcid)
    note = d.add_paragraph().add_run(
        f"Prepared {timezone.localdate():%B %-d, %Y}. Verify all entries against the sponsor's current format."
    )
    note.italic = True
    note.font.size = Pt(9)

    for heading, items in (("Active", active), ("Pending", pending)):
        d.add_paragraph().add_run(heading).bold = True
        if not items:
            d.add_paragraph("None")
        for item in items:
            rows = [
                ("Title", item["title"]),
                ("Major goals", item["goals"] or ""),
                ("Status of support", item["status_label"]),
                ("Project number", item["number"] or ""),
                ("Name of PD/PI", item["pi"] or ""),
                ("Source of support", item["source"]),
                ("Primary place of performance", item["place"] or ""),
                ("Project/proposal start and end date (MM/YYYY)", f"{_fmt_date(item['start'])} – {_fmt_date(item['end'])}"),
                ("Total award amount (including indirect costs)", _fmt_money(item["amount"])),
                ("Person months (calendar) per budget period", _pm_text(item["pm_years"])),
            ]
            table = d.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for label, value in rows:
                cells = table.add_row().cells
                cells[0].text = label
                cells[0].paragraphs[0].runs[0].bold = True
                cells[1].text = str(value)
            d.add_paragraph()
    d.add_paragraph().add_run("Overlap").bold = True
    d.add_paragraph("Describe any scientific, budgetary or commitment overlap here.")
    buf = io.BytesIO()
    d.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    resp["Content-Disposition"] = f'attachment; filename="current-pending-{person.last_name.lower()}-{timezone.localdate():%Y%m%d}.docx"'
    return resp
